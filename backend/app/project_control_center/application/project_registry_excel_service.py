from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from ..domain.models import ProjectDefinition

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 500
EXPECTED_HEADERS = (
    "project_id",
    "name",
    "description",
    "owner",
    "active",
    "created_at",
)


class ProjectRegistryExcelError(ValueError):
    """Raised for safe, user-correctable Excel import errors."""


@dataclass(frozen=True, slots=True)
class ProjectRegistryImportResult:
    imported: int
    mode: str
    projects: tuple[ProjectDefinition, ...]


class ProjectRegistryExcelService:
    """Validated Excel interchange for the authoritative Project Registry."""

    def __init__(self, registry_service, bulk_writer) -> None:
        if registry_service is None:
            raise ValueError("registry_service is required.")
        if bulk_writer is None:
            raise ValueError("bulk_writer is required.")
        self._registry_service = registry_service
        self._bulk_writer = bulk_writer

    def export_projects(self) -> bytes:
        projects = self._registry_service.list_projects()
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Projects"
        worksheet.append(EXPECTED_HEADERS)

        for project in projects:
            worksheet.append(
                [
                    project.project_id,
                    project.name,
                    project.description,
                    project.owner,
                    project.active,
                    _as_utc(project.created_at).isoformat(),
                ]
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for index, width in enumerate((24, 32, 60, 32, 12, 32), start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = width

        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    def import_projects(
        self,
        content: bytes,
        *,
        mode: str = "create_only",
    ) -> ProjectRegistryImportResult:
        if mode not in {"create_only", "upsert"}:
            raise ProjectRegistryExcelError(
                "Import mode must be 'create_only' or 'upsert'."
            )
        if not isinstance(content, bytes) or not content:
            raise ProjectRegistryExcelError("Excel workbook is empty.")
        if len(content) > MAX_FILE_BYTES:
            raise ProjectRegistryExcelError("Excel workbook exceeds the 5 MB limit.")
        if not content.startswith(b"PK"):
            raise ProjectRegistryExcelError("The uploaded file is not a valid XLSX workbook.")

        self._reject_unsafe_zip(content)
        try:
            workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=False,
            )
        except (BadZipFile, OSError, ValueError) as exc:
            raise ProjectRegistryExcelError(
                "The uploaded file is not a valid XLSX workbook."
            ) from exc

        try:
            if workbook.sheetnames != ["Projects"]:
                raise ProjectRegistryExcelError(
                    "Workbook must contain exactly one sheet named 'Projects'."
                )
            worksheet = workbook["Projects"]
            rows = worksheet.iter_rows(values_only=False)
            try:
                header_cells = next(rows)
            except StopIteration as exc:
                raise ProjectRegistryExcelError("Workbook contains no header row.") from exc

            headers = tuple(
                str(cell.value).strip() if cell.value is not None else ""
                for cell in header_cells
            )
            if headers != EXPECTED_HEADERS:
                raise ProjectRegistryExcelError(
                    "Workbook headers do not match the Project Registry template."
                )

            projects: list[ProjectDefinition] = []
            seen_ids: set[str] = set()
            data_rows = 0

            for row_number, row in enumerate(rows, start=2):
                if all(cell.value is None for cell in row):
                    continue
                data_rows += 1
                if data_rows > MAX_ROWS:
                    raise ProjectRegistryExcelError(
                        f"Workbook exceeds the {MAX_ROWS} data-row limit."
                    )
                if any(cell.data_type == "f" for cell in row):
                    raise ProjectRegistryExcelError(
                        f"Formulas are not permitted in row {row_number}."
                    )

                values = [cell.value for cell in row]
                if len(values) < len(EXPECTED_HEADERS):
                    values.extend([None] * (len(EXPECTED_HEADERS) - len(values)))

                project_id = _required_text(values[0], "project_id", row_number, 128)
                name = _required_text(values[1], "name", row_number, 255)
                description = _optional_text(values[2], "description", row_number, 2000) or ""
                owner = _optional_text(values[3], "owner", row_number, 255)
                active = _boolean(values[4], "active", row_number)

                if project_id in seen_ids:
                    raise ProjectRegistryExcelError(
                        f"Duplicate project_id '{project_id}' in row {row_number}."
                    )
                seen_ids.add(project_id)

                projects.append(
                    ProjectDefinition(
                        project_id=project_id,
                        name=name,
                        description=description,
                        owner=owner,
                        active=active,
                    )
                )

            if not projects:
                raise ProjectRegistryExcelError("Workbook contains no project rows.")

            if mode == "create_only":
                existing = [
                    project.project_id
                    for project in projects
                    if self._registry_service.get_project(project.project_id) is not None
                ]
                if existing:
                    raise ProjectRegistryExcelError(
                        "Project IDs already exist: " + ", ".join(sorted(existing))
                    )
                persisted = self._bulk_writer.create_many(projects)
            else:
                persisted = self._bulk_writer.upsert_many(projects)

            return ProjectRegistryImportResult(
                imported=len(persisted),
                mode=mode,
                projects=tuple(persisted),
            )
        finally:
            workbook.close()

    @staticmethod
    def _reject_unsafe_zip(content: bytes) -> None:
        try:
            with ZipFile(BytesIO(content)) as archive:
                names = archive.namelist()
                if any(name.lower().endswith("vbaproject.bin") for name in names):
                    raise ProjectRegistryExcelError(
                        "Macro-enabled workbooks are not permitted."
                    )
                if len(names) > 2000:
                    raise ProjectRegistryExcelError(
                        "Workbook contains too many archive entries."
                    )
        except BadZipFile as exc:
            raise ProjectRegistryExcelError(
                "The uploaded file is not a valid XLSX workbook."
            ) from exc


def _required_text(value, field: str, row: int, max_length: int) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ProjectRegistryExcelError(f"{field} is required in row {row}.")
    result = value.strip()
    if len(result) > max_length:
        raise ProjectRegistryExcelError(
            f"{field} exceeds {max_length} characters in row {row}."
        )
    return result


def _optional_text(value, field: str, row: int, max_length: int) -> str | None:
    if value is None:
        return None
    if not isinstance(value, str):
        raise ProjectRegistryExcelError(f"{field} must be text in row {row}.")
    result = value.strip()
    if len(result) > max_length:
        raise ProjectRegistryExcelError(
            f"{field} exceeds {max_length} characters in row {row}."
        )
    return result or None


def _boolean(value, field: str, row: int) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in {"true", "yes", "1"}:
            return True
        if normalized in {"false", "no", "0"}:
            return False
    if isinstance(value, (int, float)) and value in {0, 1}:
        return bool(value)
    raise ProjectRegistryExcelError(f"{field} must be boolean in row {row}.")


def _as_utc(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


__all__ = [
    "EXPECTED_HEADERS",
    "MAX_FILE_BYTES",
    "MAX_ROWS",
    "ProjectRegistryExcelError",
    "ProjectRegistryExcelService",
    "ProjectRegistryImportResult",
]

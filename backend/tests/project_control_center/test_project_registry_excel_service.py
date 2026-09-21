from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.project_control_center.application.project_registry_excel_service import (
    EXPECTED_HEADERS,
    MAX_ROWS,
    ProjectRegistryExcelError,
    ProjectRegistryExcelService,
)
from app.project_control_center.domain.models import ProjectDefinition


class FakeRegistryService:
    def __init__(self, projects: list[ProjectDefinition] | None = None) -> None:
        self.projects = {item.project_id: item for item in projects or []}

    def list_projects(self) -> list[ProjectDefinition]:
        return list(self.projects.values())

    def get_project(self, project_id: str) -> ProjectDefinition | None:
        return self.projects.get(project_id)


class FakeBulkWriter:
    def __init__(self, registry: FakeRegistryService) -> None:
        self.registry = registry
        self.created: list[ProjectDefinition] = []
        self.upserted: list[ProjectDefinition] = []

    def create_many(self, projects: list[ProjectDefinition]) -> list[ProjectDefinition]:
        self.created.extend(projects)
        for project in projects:
            self.registry.projects[project.project_id] = project
        return projects

    def upsert_many(self, projects: list[ProjectDefinition]) -> list[ProjectDefinition]:
        self.upserted.extend(projects)
        for project in projects:
            self.registry.projects[project.project_id] = project
        return projects


def make_workbook(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Projects"
    sheet.append(list(EXPECTED_HEADERS))
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def make_service(
    projects: list[ProjectDefinition] | None = None,
) -> tuple[ProjectRegistryExcelService, FakeRegistryService, FakeBulkWriter]:
    registry = FakeRegistryService(projects)
    writer = FakeBulkWriter(registry)
    return ProjectRegistryExcelService(registry, writer), registry, writer


def test_export_contains_expected_contract() -> None:
    project = ProjectDefinition(
        project_id="easychatbot",
        name="EasyChatBot",
        description="Authoritative project",
        owner="Architecture",
        active=True,
    )
    service, _, _ = make_service([project])

    workbook = load_workbook(BytesIO(service.export_projects()), read_only=True, data_only=False)
    sheet = workbook["Projects"]

    assert tuple(cell.value for cell in next(sheet.iter_rows())) == EXPECTED_HEADERS
    values = [cell.value for cell in next(sheet.iter_rows(min_row=2))]
    assert values[:5] == ["easychatbot", "EasyChatBot", "Authoritative project", "Architecture", True]
    workbook.close()


def test_import_create_only_is_validated_before_persistence() -> None:
    service, _, writer = make_service()
    content = make_workbook([["new-project", "New Project", "desc", "owner", True, "ignored"]])

    result = service.import_projects(content)

    assert result.imported == 1
    assert result.mode == "create_only"
    assert writer.created[0].project_id == "new-project"
    assert writer.created[0].created_at.tzinfo is not None


def test_import_rejects_duplicate_ids_inside_workbook() -> None:
    service, _, writer = make_service()
    content = make_workbook([
        ["duplicate", "One", "", "", True, ""],
        ["duplicate", "Two", "", "", True, ""],
    ])

    with pytest.raises(ProjectRegistryExcelError, match="Duplicate project_id"):
        service.import_projects(content)

    assert writer.created == []


def test_import_create_only_rejects_existing_project() -> None:
    existing = ProjectDefinition(project_id="existing", name="Existing")
    service, _, writer = make_service([existing])
    content = make_workbook([["existing", "Replacement", "", "", True, ""]])

    with pytest.raises(ProjectRegistryExcelError, match="already exist"):
        service.import_projects(content)

    assert writer.created == []


def test_import_upsert_updates_existing_project() -> None:
    existing = ProjectDefinition(project_id="existing", name="Existing")
    service, registry, writer = make_service([existing])
    content = make_workbook([["existing", "Updated", "new", "owner", False, ""]])

    result = service.import_projects(content, mode="upsert")

    assert result.imported == 1
    assert writer.upserted[0].name == "Updated"
    assert registry.projects["existing"].active is False


def test_import_rejects_formulas() -> None:
    service, _, writer = make_service()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Projects"
    sheet.append(list(EXPECTED_HEADERS))
    sheet.append(["formula-project", "=1+1", "", "", True, ""])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(ProjectRegistryExcelError, match="Formulas are not permitted"):
        service.import_projects(output.getvalue())

    assert writer.created == []


def test_import_rejects_excess_rows() -> None:
    service, _, writer = make_service()
    rows = [
        [f"project-{index}", f"Project {index}", "", "", True, ""]
        for index in range(MAX_ROWS + 1)
    ]

    with pytest.raises(ProjectRegistryExcelError, match="data-row limit"):
        service.import_projects(make_workbook(rows))

    assert writer.created == []

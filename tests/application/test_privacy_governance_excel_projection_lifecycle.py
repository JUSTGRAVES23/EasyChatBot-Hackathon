from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from app.application.privacy.privacy_approval_record import (
    ApprovalDecision,
    PrivacyApprovalRecord,
)
from app.application.privacy.privacy_approval_registry import (
    PrivacyApprovalRegistry,
)
from app.application.privacy.privacy_governance_case import (
    PrivacyGovernanceCase,
    PrivacyTriggerType,
)
from app.application.privacy.privacy_governance_case_registry import (
    PrivacyGovernanceCaseRegistry,
)
from app.application.privacy.privacy_governance_excel_sync import (
    PrivacyGovernanceExcelSync,
)


def make_case(status: str) -> PrivacyGovernanceCase:
    return PrivacyGovernanceCase(
        case_id="CASE-LIFECYCLE-001",
        title="Lifecycle test",
        summary="Lifecycle projection.",
        trigger_type=PrivacyTriggerType.NEW_PROCESSING,
        triggered_at=datetime.now(timezone.utc),
        trigger_source="Processing Registry",
        risk="MEDIUM",
        validation_status="PASS",
        email_status="SENT",
        governance_status=status,
    )


def make_approval(
    approval_id: str,
    role: str,
    decision: ApprovalDecision,
) -> PrivacyApprovalRecord:
    return PrivacyApprovalRecord(
        approval_id=approval_id,
        case_id="CASE-LIFECYCLE-001",
        change_id="PCR-LIFECYCLE-001",
        reviewer=f"{role} Reviewer",
        role=role,
        decision=decision,
        comment="Lifecycle decision.",
        decided_at=datetime.now(timezone.utc),
    )


def read_dashboard(path: Path) -> dict[str, int]:
    ws = load_workbook(path)["Governance Dashboard"]
    return {
        ws[f"A{row}"].value: ws[f"B{row}"].value
        for row in range(2, ws.max_row + 1)
    }


def test_dpo_plus_grc_projects_approved_state(tmp_path: Path):
    cases = PrivacyGovernanceCaseRegistry([make_case("APPROVED")])
    approvals = PrivacyApprovalRegistry()
    approvals.register(
        make_approval("APR-DPO", "DPO", ApprovalDecision.APPROVED)
    )
    approvals.register(
        make_approval("APR-GRC", "GRC", ApprovalDecision.APPROVED)
    )

    sync = PrivacyGovernanceExcelSync(
        workbook_path=tmp_path / "tracker.xlsx"
    )
    sync.sync_all(
        case_registry=cases,
        approval_registry=approvals,
    )

    dashboard = read_dashboard(tmp_path / "tracker.xlsx")

    assert dashboard["Total Cases"] == 1
    assert dashboard["Pending DPO/GRC"] == 0
    assert dashboard["Approved"] == 1
    assert dashboard["Blocked"] == 0
    assert dashboard["Approval Records"] == 2


def test_rejection_projects_blocked_state(tmp_path: Path):
    cases = PrivacyGovernanceCaseRegistry([make_case("BLOCKED")])
    approvals = PrivacyApprovalRegistry()
    approvals.register(
        make_approval(
            "APR-DPO",
            "DPO",
            ApprovalDecision.REJECTED,
        )
    )

    sync = PrivacyGovernanceExcelSync(
        workbook_path=tmp_path / "tracker.xlsx"
    )
    sync.sync_all(
        case_registry=cases,
        approval_registry=approvals,
    )

    dashboard = read_dashboard(tmp_path / "tracker.xlsx")

    assert dashboard["Total Cases"] == 1
    assert dashboard["Pending DPO/GRC"] == 0
    assert dashboard["Approved"] == 0
    assert dashboard["Blocked"] == 1
    assert dashboard["Approval Records"] == 1


def test_request_changes_projects_blocked_state(tmp_path: Path):
    cases = PrivacyGovernanceCaseRegistry([make_case("BLOCKED")])
    approvals = PrivacyApprovalRegistry()
    approvals.register(
        make_approval(
            "APR-GRC",
            "GRC",
            ApprovalDecision.REQUEST_CHANGES,
        )
    )

    sync = PrivacyGovernanceExcelSync(
        workbook_path=tmp_path / "tracker.xlsx"
    )
    sync.sync_all(
        case_registry=cases,
        approval_registry=approvals,
    )

    dashboard = read_dashboard(tmp_path / "tracker.xlsx")

    assert dashboard["Total Cases"] == 1
    assert dashboard["Pending DPO/GRC"] == 0
    assert dashboard["Approved"] == 0
    assert dashboard["Blocked"] == 1
    assert dashboard["Approval Records"] == 1


from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from backend.app.application.privacy.privacy_approval_record import (
    ApprovalDecision,
)
from backend.app.application.privacy.privacy_change import (
    PrivacyChange,
    PrivacyChangeRisk,
)
from backend.app.application.privacy.privacy_governance_composition import (
    create_privacy_governance_composition,
)
from backend.app.application.privacy.privacy_governance_excel_sync import (
    PrivacyGovernanceExcelSync,
)
from backend.app.application.privacy.privacy_governance_runtime_sync import (
    PrivacyGovernanceRuntimeSync,
)


def read_sheet_rows(
    workbook_path: Path,
    sheet_name: str,
) -> list[dict[str, object]]:
    workbook = load_workbook(workbook_path)

    worksheet = workbook[sheet_name]

    rows = list(
        worksheet.iter_rows(
            values_only=True
        )
    )

    if not rows:
        return []

    headers = rows[0]

    return [
        dict(zip(headers, row))
        for row in rows[1:]
    ]


def read_dashboard(
    workbook_path: Path,
) -> dict[str, object]:
    workbook = load_workbook(workbook_path)

    worksheet = workbook[
        "Governance Dashboard"
    ]

    return {
        worksheet[f"A{row}"].value:
        worksheet[f"B{row}"].value
        for row in range(
            2,
            worksheet.max_row + 1,
        )
    }


def make_change() -> PrivacyChange:
    return PrivacyChange(
        change_id="PCR-6H-E2E-001",
        title="6H Full Lifecycle Processing",
        description=(
            "End-to-end privacy governance Excel "
            "integration test."
        ),
        risk=PrivacyChangeRisk.MEDIUM,
        dpo_required=True,
        grc_required=True,
    )


def test_full_approved_lifecycle_projects_to_excel(
    tmp_path: Path,
):
    """
    Prove the complete 6H lifecycle:

        NEW_PROCESSING
            ↓
        Validation PASS
            ↓
        Governance Case
            ↓
        PENDING_DPO_GRC
            ↓
        DPO approval
            ↓
        GRC approval
            ↓
        APPROVED
            ↓
        Excel evidence

    The application registries remain authoritative.
    Excel is only a downstream evidence projection.
    """

    composition = (
        create_privacy_governance_composition()
    )

    workbook_path = (
        tmp_path
        / "privacy_governance_tracker.xlsx"
    )

    excel_sync = PrivacyGovernanceExcelSync(
        workbook_path=workbook_path,
    )

    runtime_sync = PrivacyGovernanceRuntimeSync(
        case_service=composition.case_service,
        excel_sync=excel_sync,
    )

    # ---------------------------------------------------------
    # 1. Real NEW_PROCESSING enters the governance flow.
    # ---------------------------------------------------------

    result = (
        composition.flow_service
        .process_new_processing(
            make_change(),
            automated_validation_passed=True,
            case_id="CASE-6H-E2E-001",
            trigger_source="Processing Registry",
        )
    )

    assert (
        result.case.case_id
        == "CASE-6H-E2E-001"
    )

    assert (
        result.case.change_id
        == "PCR-6H-E2E-001"
    )

    # ---------------------------------------------------------
    # 2. Validation PASS.
    # ---------------------------------------------------------

    assert (
        result.case.validation_status
        == "PASS"
    )

    # ---------------------------------------------------------
    # 3. Governance starts pending.
    # ---------------------------------------------------------

    assert (
        result.case.governance_status
        == "PENDING_DPO_GRC"
    )

    # ---------------------------------------------------------
    # 4. Runtime case sync → Excel.
    # ---------------------------------------------------------

    synced_case = runtime_sync.sync_case(
        "CASE-6H-E2E-001"
    )

    assert (
        synced_case.case_id
        == "CASE-6H-E2E-001"
    )

    assert (
        synced_case.change_id
        == "PCR-6H-E2E-001"
    )

    # ---------------------------------------------------------
    # 5. Verify initial Excel case evidence.
    # ---------------------------------------------------------

    case_rows = read_sheet_rows(
        workbook_path,
        "Privacy Governance Cases",
    )

    assert len(case_rows) == 1

    case_row = case_rows[0]

    assert (
        case_row["Case ID"]
        == "CASE-6H-E2E-001"
    )

    assert (
        case_row["Change ID"]
        == "PCR-6H-E2E-001"
    )

    assert (
        case_row["Title"]
        == "6H Full Lifecycle Processing"
    )

    assert (
        case_row["Risk"]
        == "MEDIUM"
    )

    assert (
        case_row["Validation"]
        == "PASS"
    )

    assert (
        case_row["Governance Status"]
        == "PENDING_DPO_GRC"
    )

    # ---------------------------------------------------------
    # 6. DPO approval.
    # ---------------------------------------------------------

    dpo_record = (
        composition.case_service.approve(
            case_id="CASE-6H-E2E-001",
            change_id="PCR-6H-E2E-001",
            approval_id="APR-6H-DPO-001",
            reviewer="DPO Reviewer",
            role="DPO",
            comment=(
                "DPO approval for 6H E2E test."
            ),
        )
    )

    assert (
        dpo_record.decision
        == ApprovalDecision.APPROVED
    )

    assert (
        composition.case_registry
        .get("CASE-6H-E2E-001")
        .governance_status
        == "PENDING_DPO_GRC"
    )

    # ---------------------------------------------------------
    # 7. Sync DPO decision to Excel.
    # ---------------------------------------------------------

    runtime_sync.sync_decision(
        "CASE-6H-E2E-001"
    )

    approval_rows = read_sheet_rows(
        workbook_path,
        "Approval Evidence",
    )

    assert len(approval_rows) == 1

    dpo_row = approval_rows[0]

    assert (
        dpo_row["Approval ID"]
        == "APR-6H-DPO-001"
    )

    assert (
        dpo_row["Case ID"]
        == "CASE-6H-E2E-001"
    )

    assert (
        dpo_row["Change ID"]
        == "PCR-6H-E2E-001"
    )

    assert (
        dpo_row["Reviewer"]
        == "DPO Reviewer"
    )

    assert (
        dpo_row["Role"]
        == "DPO"
    )

    assert (
        dpo_row["Decision"]
        == "APPROVED"
    )

    # ---------------------------------------------------------
    # 8. GRC approval.
    # ---------------------------------------------------------

    grc_record = (
        composition.case_service.approve(
            case_id="CASE-6H-E2E-001",
            change_id="PCR-6H-E2E-001",
            approval_id="APR-6H-GRC-001",
            reviewer="GRC Reviewer",
            role="GRC",
            comment=(
                "GRC approval for 6H E2E test."
            ),
        )
    )

    assert (
        grc_record.decision
        == ApprovalDecision.APPROVED
    )

    # ---------------------------------------------------------
    # 9. Final governance state.
    # ---------------------------------------------------------

    final_case = (
        composition.case_registry
        .get("CASE-6H-E2E-001")
    )

    final_change = (
        composition.change_registry
        .get("PCR-6H-E2E-001")
    )

    assert (
        final_case.governance_status
        == "APPROVED"
    )

    assert (
        final_case.final_decision
        == "APPROVED"
    )

    assert final_case.final_decision_at is not None

    assert final_change is not None

    assert (
        final_change.status.value
        == "approved"
    )

    # ---------------------------------------------------------
    # 10. Sync final decision to Excel.
    # ---------------------------------------------------------

    runtime_sync.sync_decision(
        "CASE-6H-E2E-001"
    )

    # ---------------------------------------------------------
    # 11. Verify final case projection.
    # ---------------------------------------------------------

    case_rows = read_sheet_rows(
        workbook_path,
        "Privacy Governance Cases",
    )

    assert len(case_rows) == 1

    final_case_row = case_rows[0]

    assert (
        final_case_row["Case ID"]
        == "CASE-6H-E2E-001"
    )

    assert (
        final_case_row["Change ID"]
        == "PCR-6H-E2E-001"
    )

    assert (
        final_case_row["Validation"]
        == "PASS"
    )

    assert (
        final_case_row["Governance Status"]
        == "APPROVED"
    )

    assert (
        final_case_row["Final Decision"]
        == "APPROVED"
    )

    # ---------------------------------------------------------
    # 12. Verify both approval records.
    # ---------------------------------------------------------

    approval_rows = read_sheet_rows(
        workbook_path,
        "Approval Evidence",
    )

    assert len(approval_rows) == 2

    approval_ids = {
        row["Approval ID"]
        for row in approval_rows
    }

    assert approval_ids == {
        "APR-6H-DPO-001",
        "APR-6H-GRC-001",
    }

    for row in approval_rows:
        assert (
            row["Case ID"]
            == "CASE-6H-E2E-001"
        )

        assert (
            row["Change ID"]
            == "PCR-6H-E2E-001"
        )

        assert (
            row["Decision"]
            == "APPROVED"
        )

    # ---------------------------------------------------------
    # 13. Verify Governance Dashboard.
    # ---------------------------------------------------------

    dashboard = read_dashboard(
        workbook_path
    )

    assert (
        dashboard["Total Cases"]
        == 1
    )

    assert (
        dashboard["Pending DPO/GRC"]
        == 0
    )

    assert (
        dashboard["Approved"]
        == 1
    )

    assert (
        dashboard["Blocked"]
        == 0
    )

    assert (
        dashboard["Approval Records"]
        == 2
    )


def test_rejection_lifecycle_projects_blocked_to_excel(
    tmp_path: Path,
):
    """
    Prove:

        Validation PASS
            ↓
        PENDING_DPO_GRC
            ↓
        Human rejection
            ↓
        BLOCKED
            ↓
        Excel evidence
    """

    composition = (
        create_privacy_governance_composition()
    )

    workbook_path = (
        tmp_path
        / "privacy_governance_rejection.xlsx"
    )

    excel_sync = PrivacyGovernanceExcelSync(
        workbook_path=workbook_path,
    )

    runtime_sync = PrivacyGovernanceRuntimeSync(
        case_service=composition.case_service,
        excel_sync=excel_sync,
    )

    result = (
        composition.flow_service
        .process_new_processing(
            PrivacyChange(
                change_id="PCR-6H-E2E-REJECT-001",
                title="Rejected processing",
                description=(
                    "6H rejection lifecycle."
                ),
                risk=PrivacyChangeRisk.MEDIUM,
                dpo_required=True,
                grc_required=True,
            ),
            automated_validation_passed=True,
            case_id="CASE-6H-E2E-REJECT-001",
            trigger_source="Processing Registry",
        )
    )

    assert (
        result.case.governance_status
        == "PENDING_DPO_GRC"
    )

    runtime_sync.sync_case(
        "CASE-6H-E2E-REJECT-001"
    )

    record = (
        composition.case_service.reject(
            case_id="CASE-6H-E2E-REJECT-001",
            change_id="PCR-6H-E2E-REJECT-001",
            approval_id="APR-6H-REJECT-001",
            reviewer="DPO Reviewer",
            role="DPO",
            comment=(
                "Rejected during 6H E2E test."
            ),
        )
    )

    assert (
        record.decision
        == ApprovalDecision.REJECTED
    )

    case = (
        composition.case_registry
        .get("CASE-6H-E2E-REJECT-001")
    )

    assert (
        case.governance_status
        == "BLOCKED"
    )

    assert (
        case.final_decision
        == "REJECTED"
    )

    runtime_sync.sync_decision(
        "CASE-6H-E2E-REJECT-001"
    )

    case_rows = read_sheet_rows(
        workbook_path,
        "Privacy Governance Cases",
    )

    assert len(case_rows) == 1

    assert (
        case_rows[0]["Governance Status"]
        == "BLOCKED"
    )

    assert (
        case_rows[0]["Final Decision"]
        == "REJECTED"
    )

    approval_rows = read_sheet_rows(
        workbook_path,
        "Approval Evidence",
    )

    assert len(approval_rows) == 1

    assert (
        approval_rows[0]["Approval ID"]
        == "APR-6H-REJECT-001"
    )

    assert (
        approval_rows[0]["Decision"]
        == "REJECTED"
    )

    dashboard = read_dashboard(
        workbook_path
    )

    assert dashboard["Total Cases"] == 1
    assert dashboard["Pending DPO/GRC"] == 0
    assert dashboard["Approved"] == 0
    assert dashboard["Blocked"] == 1
    assert dashboard["Approval Records"] == 1


def test_request_changes_lifecycle_projects_blocked_to_excel(
    tmp_path: Path,
):
    """
    Prove:

        Validation PASS
            ↓
        PENDING_DPO_GRC
            ↓
        Request changes
            ↓
        BLOCKED
            ↓
        Excel evidence
    """

    composition = (
        create_privacy_governance_composition()
    )

    workbook_path = (
        tmp_path
        / "privacy_governance_request_changes.xlsx"
    )

    excel_sync = PrivacyGovernanceExcelSync(
        workbook_path=workbook_path,
    )

    runtime_sync = PrivacyGovernanceRuntimeSync(
        case_service=composition.case_service,
        excel_sync=excel_sync,
    )

    result = (
        composition.flow_service
        .process_new_processing(
            PrivacyChange(
                change_id=(
                    "PCR-6H-E2E-CHANGES-001"
                ),
                title="Changes requested",
                description=(
                    "6H request-changes lifecycle."
                ),
                risk=PrivacyChangeRisk.MEDIUM,
                dpo_required=True,
                grc_required=True,
            ),
            automated_validation_passed=True,
            case_id=(
                "CASE-6H-E2E-CHANGES-001"
            ),
            trigger_source="Processing Registry",
        )
    )

    assert (
        result.case.governance_status
        == "PENDING_DPO_GRC"
    )

    runtime_sync.sync_case(
        "CASE-6H-E2E-CHANGES-001"
    )

    record = (
        composition.case_service
        .request_changes(
            case_id=(
                "CASE-6H-E2E-CHANGES-001"
            ),
            change_id=(
                "PCR-6H-E2E-CHANGES-001"
            ),
            approval_id=(
                "APR-6H-CHANGES-001"
            ),
            reviewer="GRC Reviewer",
            role="GRC",
            comment=(
                "Additional privacy evidence "
                "required."
            ),
        )
    )

    assert (
        record.decision
        == ApprovalDecision.REQUEST_CHANGES
    )

    case = (
        composition.case_registry
        .get(
            "CASE-6H-E2E-CHANGES-001"
        )
    )

    assert (
        case.governance_status
        == "BLOCKED"
    )

    assert (
        case.final_decision
        == "REQUEST_CHANGES"
    )

    runtime_sync.sync_decision(
        "CASE-6H-E2E-CHANGES-001"
    )

    case_rows = read_sheet_rows(
        workbook_path,
        "Privacy Governance Cases",
    )

    assert len(case_rows) == 1

    assert (
        case_rows[0]["Governance Status"]
        == "BLOCKED"
    )

    assert (
        case_rows[0]["Final Decision"]
        == "REQUEST_CHANGES"
    )

    approval_rows = read_sheet_rows(
        workbook_path,
        "Approval Evidence",
    )

    assert len(approval_rows) == 1

    assert (
        approval_rows[0]["Approval ID"]
        == "APR-6H-CHANGES-001"
    )

    assert (
        approval_rows[0]["Decision"]
        == "REQUEST_CHANGES"
    )

    dashboard = read_dashboard(
        workbook_path
    )

    assert dashboard["Total Cases"] == 1
    assert dashboard["Pending DPO/GRC"] == 0
    assert dashboard["Approved"] == 0
    assert dashboard["Blocked"] == 1
    assert dashboard["Approval Records"] == 1
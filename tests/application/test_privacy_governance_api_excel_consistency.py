from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from backend.app.application.privacy.privacy_governance_excel_sync import (
    PrivacyGovernanceExcelSync,
)


def test_excel_projection_contains_authoritative_case_and_approval_evidence(
    tmp_path: Path,
):
    # This is a focused adapter-level smoke test for the final projection.
    # Full API lifecycle tests are in tests/api.
    workbook_path = tmp_path / "tracker.xlsx"

    sync = PrivacyGovernanceExcelSync(
        workbook_path=workbook_path
    )

    assert not workbook_path.exists()

    # Empty registries are a valid initial projection.
    class EmptyCases:
        @staticmethod
        def all():
            return ()

    class EmptyApprovals:
        @staticmethod
        def for_case(case_id):
            return ()

    sync.sync_all(
        case_registry=EmptyCases(),
        approval_registry=EmptyApprovals(),
    )

    assert workbook_path.exists()

    workbook = load_workbook(workbook_path)

    assert "Privacy Governance Cases" in workbook.sheetnames
    assert "Approval Evidence" in workbook.sheetnames
    assert "Governance Dashboard" in workbook.sheetnames

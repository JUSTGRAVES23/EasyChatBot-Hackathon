from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from backend.app.application.privacy.privacy_approval_record import (
    ApprovalDecision,
    PrivacyApprovalRecord,
)
from backend.app.application.privacy.privacy_approval_registry import (
    PrivacyApprovalRegistry,
)
from backend.app.application.privacy.privacy_governance_case import (
    PrivacyGovernanceCase,
    PrivacyTriggerType,
)
from backend.app.application.privacy.privacy_governance_case_registry import (
    PrivacyGovernanceCaseRegistry,
)
from backend.app.application.privacy.privacy_governance_excel_sync import (
    PrivacyGovernanceExcelSync,
)


def make_case(
    *,
    case_id: str = "CASE-001",
    status: str = "PENDING_DPO_GRC",
) -> PrivacyGovernanceCase:
    return PrivacyGovernanceCase(
        case_id=case_id,
        title="New learner processing",
        summary="Controlled privacy governance case.",
        trigger_type=PrivacyTriggerType.NEW_PROCESSING,
        triggered_at=datetime.now(timezone.utc),
        trigger_source="Processing Registry",
        risk="MEDIUM",
        validation_status="PASS",
        email_status="SENT",
        governance_status=status,
    )


def make_approval(
    approval_id: str,
    *,
    case_id: str = "CASE-001",
    role: str = "DPO",
    decision: ApprovalDecision = ApprovalDecision.APPROVED,
) -> PrivacyApprovalRecord:
    return PrivacyApprovalRecord(
        approval_id=approval_id,
        case_id=case_id,
        change_id="PCR-001",
        reviewer=f"{role} Reviewer",
        role=role,
        decision=decision,
        comment=f"{role} decision recorded.",
        decided_at=datetime.now(timezone.utc),
    )


def make_sync(tmp_path: Path) -> PrivacyGovernanceExcelSync:
    return PrivacyGovernanceExcelSync(
        workbook_path=tmp_path / "EasyChatBot_Project_Control_Center.xlsx",
    )


def test_case_is_projected_to_excel(tmp_path: Path):
    cases = PrivacyGovernanceCaseRegistry([make_case()])
    approvals = PrivacyApprovalRegistry()

    make_sync(tmp_path).sync_all(
        case_registry=cases,
        approval_registry=approvals,
    )

    workbook = load_workbook(
        tmp_path / "EasyChatBot_Project_Control_Center.xlsx"
    )
    ws = workbook["Privacy Governance Cases"]

    assert ws["A2"].value == "CASE-001"
    assert ws["I2"].value == "PENDING_DPO_GRC"


def test_approval_is_projected_to_excel(tmp_path: Path):
    cases = PrivacyGovernanceCaseRegistry([make_case()])
    approvals = PrivacyApprovalRegistry()
    approvals.register(make_approval("APR-DPO-001"))

    make_sync(tmp_path).sync_all(
        case_registry=cases,
        approval_registry=approvals,
    )

    workbook = load_workbook(
        tmp_path / "EasyChatBot_Project_Control_Center.xlsx"
    )
    ws = workbook["Approval Evidence"]

    assert ws["A2"].value == "APR-DPO-001"
    assert ws["B2"].value == "CASE-001"
    assert ws["E2"].value == "DPO"
    assert ws["F2"].value == "APPROVED"


def test_dashboard_counts_are_projected(tmp_path: Path):
    cases = PrivacyGovernanceCaseRegistry([
        make_case(case_id="CASE-001", status="APPROVED"),
        make_case(case_id="CASE-002", status="BLOCKED"),
        make_case(case_id="CASE-003", status="PENDING_DPO_GRC"),
    ])
    approvals = PrivacyApprovalRegistry()
    approvals.register(make_approval("APR-001", case_id="CASE-001"))
    approvals.register(
        make_approval(
            "APR-002",
            case_id="CASE-001",
            role="GRC",
        )
    )

    make_sync(tmp_path).sync_all(
        case_registry=cases,
        approval_registry=approvals,
    )

    workbook = load_workbook(
        tmp_path / "EasyChatBot_Project_Control_Center.xlsx"
    )
    ws = workbook["Governance Dashboard"]

    values = {
        ws[f"A{row}"].value: ws[f"B{row}"].value
        for row in range(2, ws.max_row + 1)
    }

    assert values["Total Cases"] == 3
    assert values["Pending DPO/GRC"] == 1
    assert values["Approved"] == 1
    assert values["Blocked"] == 1
    assert values["Approval Records"] == 2


def test_existing_workbook_sheets_are_preserved(tmp_path: Path):
    from openpyxl import Workbook

    workbook_path = tmp_path / "EasyChatBot_Project_Control_Center.xlsx"

    workbook = Workbook()
    workbook.active.title = "Beta Tracker"
    workbook.create_sheet("Capability Tracker")
    workbook.save(workbook_path)

    cases = PrivacyGovernanceCaseRegistry([make_case()])
    approvals = PrivacyApprovalRegistry()

    make_sync(tmp_path).sync_all(
        case_registry=cases,
        approval_registry=approvals,
    )

    workbook = load_workbook(workbook_path)

    assert "Beta Tracker" in workbook.sheetnames
    assert "Capability Tracker" in workbook.sheetnames
    assert "Privacy Governance Cases" in workbook.sheetnames
    assert "Approval Evidence" in workbook.sheetnames
    assert "Governance Dashboard" in workbook.sheetnames
